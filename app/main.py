"""
AcquireFlow – Flask backend
Pure Python. No Rust. Works on Python 3.9 – 3.14+.
"""

import csv
import io
import json
import logging
import math
import os
import re
import sqlite3
import threading
import time
import uuid
from datetime import datetime

import requests
from flask import Flask, Response, jsonify, request, send_from_directory
from flask_cors import CORS
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

load_dotenv()

logging.basicConfig(level=logging.INFO, format="%(levelname)s  %(message)s")
logger = logging.getLogger(__name__)

BASE_DIR   = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
STATIC_DIR = os.path.join(BASE_DIR, "static")
DB_PATH    = os.path.join(BASE_DIR, "data", "acquireflow.db")

CH_BASE      = "https://api.company-information.service.gov.uk"
POSTCODES_IO = "https://api.postcodes.io"

app = Flask(__name__, static_folder=STATIC_DIR)
CORS(app)

# ── Database ────────────────────────────────────────────────────────────────────

_db_lock = threading.Lock()


def _conn():
    os.makedirs(os.path.dirname(DB_PATH), exist_ok=True)
    c = sqlite3.connect(DB_PATH, check_same_thread=False, timeout=30)
    c.row_factory = sqlite3.Row
    return c


def init_db():
    with _conn() as c:
        c.execute("""CREATE TABLE IF NOT EXISTS searches (
            id                 TEXT PRIMARY KEY,
            status             TEXT DEFAULT 'pending',
            progress           INTEGER DEFAULT 0,
            total_found        INTEGER DEFAULT 0,
            processed          INTEGER DEFAULT 0,
            message            TEXT DEFAULT '',
            params             TEXT,
            results            TEXT DEFAULT '[]',
            criteria_met_count INTEGER DEFAULT 0,
            created_at         TEXT,
            completed_at       TEXT,
            error              TEXT
        )""")
        c.execute("""CREATE TABLE IF NOT EXISTS company_cache (
            company_number TEXT PRIMARY KEY,
            officers       TEXT,
            cached_at      TEXT
        )""")
        c.execute("CREATE INDEX IF NOT EXISTS idx_s_created ON searches(created_at DESC)")
    logger.info("DB ready: %s", DB_PATH)


def _write(sql, params=()):
    with _db_lock:
        with _conn() as c:
            c.execute(sql, params)


def _one(sql, params=()):
    with _conn() as c:
        row = c.execute(sql, params).fetchone()
        return dict(row) if row else None


def _all(sql, params=()):
    with _conn() as c:
        return [dict(r) for r in c.execute(sql, params).fetchall()]


# ── Companies House ─────────────────────────────────────────────────────────────

def ch_get(path, params=None, retries=3):
    key = os.getenv("COMPANIES_HOUSE_API_KEY", "")
    if not key or key == "your_api_key_here":
        raise RuntimeError(
            "COMPANIES_HOUSE_API_KEY is not set. "
            "Edit .env and add your key from developer.company-information.service.gov.uk"
        )
    url = CH_BASE + path
    for attempt in range(retries):
        try:
            r = requests.get(url, params=params or {}, auth=(key, ""), timeout=20)
            if r.status_code == 429:
                wait = 60 * (attempt + 1)
                logger.warning("Rate limited – waiting %ds", wait)
                time.sleep(wait)
                continue
            if r.status_code == 404:
                return {}
            r.raise_for_status()
            return r.json()
        except requests.Timeout:
            if attempt < retries - 1:
                time.sleep(5 * (attempt + 1))
            else:
                raise
        except requests.HTTPError:
            logger.error("CH %s %s", r.status_code, url)
            return {}
    return {}


def _fetch_page(sic, start, size=100):
    data = ch_get("/advanced-search/companies", {
        "sic_codes": sic, "company_status": "active",
        "start_index": start, "size": size,
    })
    return data.get("hits", 0), data.get("items", [])


def fetch_all_companies(sic, max_results):
    companies = []
    hits, first = _fetch_page(sic, 0)
    companies.extend(first)
    target = min(max_results, hits)
    for start in range(100, target, 100):
        _, page = _fetch_page(sic, start)
        companies.extend(page)
        if len(companies) >= target:
            break
        time.sleep(0.12)
    return companies[:max_results], hits


def get_officers(company_number):
    row = _one("SELECT officers, cached_at FROM company_cache WHERE company_number=?",
               (company_number,))
    if row:
        age_h = (datetime.utcnow() - datetime.fromisoformat(row["cached_at"])).total_seconds() / 3600
        if age_h < 24:
            return json.loads(row["officers"] or "{}")
    data = ch_get(f"/company/{company_number}/officers",
                  {"items_per_page": 100, "register_view": "false"})
    with _db_lock:
        with _conn() as c:
            c.execute(
                """INSERT INTO company_cache (company_number, officers, cached_at)
                   VALUES (?,?,?)
                   ON CONFLICT(company_number) DO UPDATE
                   SET officers=excluded.officers, cached_at=excluded.cached_at""",
                (company_number, json.dumps(data), datetime.utcnow().isoformat())
            )
    return data


# ── Location ────────────────────────────────────────────────────────────────────

_PC_RE = re.compile(r"\b([A-Z]{1,2}\d[A-Z\d]?\s*\d[A-Z]{2})\b", re.IGNORECASE)


def extract_postcode(text):
    m = _PC_RE.search(text or "")
    return m.group(1).upper().strip() if m else None


def postcode_to_latlong(query):
    clean = query.strip().upper().replace(" ", "")
    try:
        r = requests.get(f"{POSTCODES_IO}/postcodes/{clean}", timeout=10)
        if r.status_code == 200:
            d = r.json().get("result", {})
            if d.get("latitude"):
                return float(d["latitude"]), float(d["longitude"])
    except Exception:
        pass
    try:
        r = requests.get(f"{POSTCODES_IO}/postcodes",
                         params={"q": query.strip(), "limit": 1}, timeout=10)
        if r.status_code == 200:
            res = r.json().get("result") or []
            if res:
                return float(res[0]["latitude"]), float(res[0]["longitude"])
    except Exception:
        pass
    return None


def bulk_postcode_lookup(postcodes):
    results = {}
    unique = list({p.upper().strip() for p in postcodes if p})
    for i in range(0, len(unique), 100):
        chunk = unique[i:i+100]
        try:
            r = requests.post(f"{POSTCODES_IO}/postcodes",
                              json={"postcodes": chunk}, timeout=15)
            if r.status_code == 200:
                for item in r.json().get("result", []):
                    q, res = item.get("query",""), item.get("result")
                    if res:
                        results[q] = (float(res["latitude"]), float(res["longitude"]))
        except Exception as e:
            logger.warning("Bulk postcode lookup: %s", e)
    return results


def haversine_miles(lat1, lng1, lat2, lng2):
    R = 3958.8
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp = math.radians(lat2 - lat1)
    dl = math.radians(lng2 - lng1)
    a = math.sin(dp/2)**2 + math.cos(p1)*math.cos(p2)*math.sin(dl/2)**2
    return R * 2 * math.atan2(math.sqrt(a), math.sqrt(1-a))


# ── Scoring ─────────────────────────────────────────────────────────────────────

def calculate_score(age_years, directors, dist, min_dir_age, min_co_age, radius):
    notes = []

    # Company age 0-3
    age_s = 0.0
    if age_years is not None:
        if age_years >= min_co_age:
            age_s = min(3.0, 1.5 + (age_years / max(min_co_age, 1)) * 0.75)
            notes.append(f"Company age {age_years:.1f}y ≥ minimum {min_co_age}y")
        else:
            age_s = (age_years / max(min_co_age, 1)) * 1.0
            notes.append(f"Company age {age_years:.1f}y below minimum {min_co_age}y")

    # Directors 0-4
    dir_s = 0.0
    active = [d for d in directors if d.get("is_active")]
    qual   = [d for d in active if d.get("meets_age_criteria")]
    if qual:
        dir_s = min(4.0, 2.5 + (len(qual) - 1) * 0.5)
        notes.append(f"{len(qual)} director(s) meet age ≥ {min_dir_age}")
    elif active:
        ages = [d["estimated_age"] for d in active if d.get("estimated_age")]
        if ages:
            closest = max(ages)
            gap = min_dir_age - closest
            dir_s = 1.0 if gap <= 5 else 0.0
            notes.append(f"Closest director age {closest} (gap {gap}y)")
        else:
            notes.append("Director ages unavailable")
    else:
        notes.append("No active directors found")

    # Proximity 0-3
    prox_s = 0.0
    if dist is not None and radius > 0 and dist <= radius:
        prox_s = max(0.0, 3.0 * (1.0 - dist / radius))
        notes.append(f"Distance {dist:.1f} miles")
    elif dist is not None:
        notes.append(f"Distance {dist:.1f} miles exceeds radius")
    else:
        notes.append("Distance unknown")

    total = round(min(10.0, max(1.0, age_s + dir_s + prox_s)), 1)
    return total, {
        "company_age_score": round(age_s, 2),
        "director_score":    round(dir_s, 2),
        "proximity_score":   round(prox_s, 2),
        "total": total, "notes": notes,
    }


# ── Search worker ───────────────────────────────────────────────────────────────

def _fmt_addr(addr):
    parts = [addr.get(k,"") for k in
             ("address_line_1","address_line_2","locality","region","postal_code","country")]
    return ", ".join(p for p in parts if p)


def _parse_directors(officers_data, min_dir_age):
    now = datetime.utcnow()
    out = []
    for item in officers_data.get("items", []):
        if "director" not in item.get("officer_role","").lower():
            continue
        is_active = not item.get("resigned_on")
        dob = item.get("date_of_birth") or {}
        by, bm = dob.get("year"), dob.get("month")
        age = None
        if by:
            age = now.year - by
            if bm and now.month < bm:
                age -= 1
            age = max(0, age)
        out.append({
            "name": item.get("name","Unknown"),
            "role": item.get("officer_role","director"),
            "birth_month": bm, "birth_year": by,
            "estimated_age": age,
            "meets_age_criteria": age is not None and age >= min_dir_age,
            "is_active": is_active,
            "appointed_on": item.get("appointed_on"),
        })
    return out


def _co_age(date_str):
    if not date_str:
        return None
    try:
        return round((datetime.utcnow() - datetime.strptime(date_str, "%Y-%m-%d")).days / 365.25, 2)
    except ValueError:
        return None


def _upd(sid, status, prog, found, proc, msg, met=0):
    _write(
        "UPDATE searches SET status=?,progress=?,total_found=?,processed=?,message=?,criteria_met_count=? WHERE id=?",
        (status, prog, found, proc, msg, met, sid)
    )


def run_search(sid, params):
    sic     = params["sic_code"]
    pc      = params["postcode"]
    radius  = float(params["radius_miles"])
    min_da  = int(params["min_director_age"])
    min_ca  = int(params["min_company_age"])
    max_r   = int(params["max_results"])

    try:
        # 1 – Resolve origin
        _upd(sid, "running", 2, 0, 0, "Resolving location…")
        origin = postcode_to_latlong(pc)
        if not origin:
            raise ValueError(f"Cannot resolve '{pc}' to a UK location. Use a full postcode e.g. SW1A 1AA")
        olat, olng = origin
        logger.info("[%s] Origin %.4f,%.4f", sid, olat, olng)

        # 2 – Fetch companies
        _upd(sid, "running", 5, 0, 0, f"Fetching companies for SIC {sic}…")
        companies, total_hits = fetch_all_companies(sic, max_r)
        total = len(companies)
        logger.info("[%s] %d companies (hits=%d)", sid, total, total_hits)
        if not total:
            _write("UPDATE searches SET status='complete',progress=100,message=? WHERE id=?",
                   ("No companies found for this SIC code.", sid))
            return

        _upd(sid, "running", 10, total, 0, f"Found {total} companies. Resolving postcodes…")

        # 3 – Bulk postcode lookup
        pc_map = {}
        to_lookup = []
        for c in companies:
            addr = c.get("registered_office_address") or {}
            p = addr.get("postal_code") or extract_postcode(_fmt_addr(addr))
            p = p.upper().strip() if p else None
            pc_map[c["company_number"]] = p
            if p:
                to_lookup.append(p)
        coords = bulk_postcode_lookup(to_lookup)

        # 4 – Radius filter
        in_radius = []
        for c in companies:
            p = pc_map.get(c["company_number"])
            if p and p in coords:
                lat, lng = coords[p]
                dist = haversine_miles(olat, olng, lat, lng)
                if dist <= radius:
                    c["_dist"] = dist
                    c["_pc"] = p
                    in_radius.append(c)
            else:
                c["_dist"] = None
                c["_pc"] = p
                in_radius.append(c)

        logger.info("[%s] %d in radius (%.1f mi)", sid, len(in_radius), radius)
        _upd(sid, "running", 20, total, 0,
             f"{len(in_radius)} companies in radius. Fetching officer data…")

        # 5 – Officers + build results
        results = []
        processed = 0
        for raw in in_radius:
            try:
                officers_data = get_officers(raw["company_number"])
            except Exception as e:
                logger.warning("Officers error %s: %s", raw["company_number"], e)
                officers_data = {}

            processed += 1
            addr_str  = _fmt_addr(raw.get("registered_office_address") or {})
            dist      = raw.get("_dist")
            inc_date  = raw.get("date_of_creation")
            age_yrs   = _co_age(inc_date)
            directors = _parse_directors(officers_data, min_da)
            active    = [d for d in directors if d["is_active"]]
            qual      = [d for d in active if d["meets_age_criteria"]]
            dir_met   = len(qual) > 0
            age_met   = age_yrs is not None and age_yrs >= min_ca
            meets_all = dir_met and age_met
            score, breakdown = calculate_score(age_yrs, directors, dist, min_da, min_ca, radius)

            results.append({
                "company_number":            raw.get("company_number",""),
                "company_name":              raw.get("company_name",""),
                "company_status":            raw.get("company_status","active"),
                "registered_address":        addr_str,
                "postcode":                  raw.get("_pc"),
                "distance_miles":            round(dist, 2) if dist is not None else None,
                "incorporation_date":        inc_date,
                "company_age_years":         age_yrs,
                "sic_codes":                 raw.get("sic_codes",[]),
                "directors":                 directors,
                "active_director_count":     len(active),
                "qualifying_director_count": len(qual),
                "director_criteria_met":     dir_met,
                "company_age_criteria_met":  age_met,
                "meets_all_criteria":        meets_all,
                "score":                     score,
                "score_breakdown":           breakdown,
            })

            if processed % 10 == 0 or processed == len(in_radius):
                pct = 20 + int((processed / max(len(in_radius), 1)) * 75)
                met_so_far = sum(1 for r in results if r["meets_all_criteria"])
                _upd(sid, "running", min(pct, 95), total, processed,
                     f"Processing {processed}/{len(in_radius)} companies…", met_so_far)

        # 6 – Sort & save
        results.sort(key=lambda r: (r["meets_all_criteria"], r["score"]), reverse=True)
        met_final = sum(1 for r in results if r["meets_all_criteria"])
        _write(
            "UPDATE searches SET status='complete',progress=100,results=?,criteria_met_count=?,completed_at=?,message=? WHERE id=?",
            (json.dumps(results), met_final, datetime.utcnow().isoformat(),
             f"Complete – {len(results)} companies, {met_final} meet criteria.", sid)
        )
        logger.info("[%s] Done. %d results, %d met criteria.", sid, len(results), met_final)

    except Exception as exc:
        logger.exception("[%s] FAILED: %s", sid, exc)
        _write("UPDATE searches SET status='error',error=?,message=? WHERE id=?",
               (str(exc), f"Error: {exc}", sid))


# ── Routes ───────────────────────────────────────────────────────────────────────

@app.route("/")
def index():
    return send_from_directory(STATIC_DIR, "index.html")

@app.route("/static/<path:filename>")
def static_files(filename):
    return send_from_directory(STATIC_DIR, filename)


@app.post("/api/search")
def start_search():
    data = request.get_json(force=True) or {}
    if not data.get("sic_code") or not data.get("postcode"):
        return jsonify({"error": "sic_code and postcode are required"}), 400
    params = {
        "sic_code":         str(data["sic_code"]).strip(),
        "postcode":         str(data["postcode"]).strip(),
        "radius_miles":     float(data.get("radius_miles", 25)),
        "min_director_age": int(data.get("min_director_age", 40)),
        "min_company_age":  int(data.get("min_company_age", 5)),
        "max_results":      int(data.get("max_results", 500)),
    }
    sid = str(uuid.uuid4())
    _write("INSERT INTO searches (id,status,params,created_at) VALUES (?,?,?,?)",
           (sid, "pending", json.dumps(params), datetime.utcnow().isoformat()))
    threading.Thread(target=run_search, args=(sid, params), daemon=True).start()
    return jsonify({"search_id": sid, "status": "pending"})


@app.get("/api/search/<sid>/status")
def get_status(sid):
    row = _one("SELECT * FROM searches WHERE id=?", (sid,))
    if not row:
        return jsonify({"error": "Not found"}), 404
    return jsonify({k: row[k] for k in
                    ("status","progress","total_found","processed","message","criteria_met_count")}
                   | {"search_id": sid})


@app.get("/api/search/<sid>/results")
def get_results(sid):
    row = _one("SELECT * FROM searches WHERE id=?", (sid,))
    if not row:
        return jsonify({"error": "Not found"}), 404
    criteria_only = request.args.get("criteria_only","false").lower() == "true"
    results = json.loads(row["results"] or "[]")
    if criteria_only:
        results = [r for r in results if r.get("meets_all_criteria")]
    params = json.loads(row["params"] or "{}")
    return jsonify({
        "search_id": sid, "status": row["status"],
        "results": results, "total_results": len(results),
        "criteria_met": sum(1 for r in results if r.get("meets_all_criteria")),
        "search_params": params, "created_at": row["created_at"],
    })


@app.get("/api/searches")
def list_searches():
    rows = _all("SELECT id,status,progress,total_found,processed,message,params,criteria_met_count,created_at,completed_at FROM searches ORDER BY created_at DESC LIMIT 50")
    out = []
    for r in rows:
        p = json.loads(r.get("params") or "{}")
        out.append({**{k: r[k] for k in ("status","progress","total_found","processed","criteria_met_count","message","created_at","completed_at")},
                    "search_id": r["id"],
                    "sic_code": p.get("sic_code",""), "postcode": p.get("postcode",""),
                    "radius_miles": p.get("radius_miles","")})
    return jsonify(out)


@app.delete("/api/search/<sid>")
def delete_search(sid):
    _write("DELETE FROM searches WHERE id=?", (sid,))
    return jsonify({"deleted": sid})


@app.get("/api/export/<sid>")
def export_results(sid):
    row = _one("SELECT * FROM searches WHERE id=?", (sid,))
    if not row:
        return jsonify({"error": "Not found"}), 404
    if row["status"] != "complete":
        return jsonify({"error": "Search not complete yet"}), 400
    criteria_only = request.args.get("criteria_only","false").lower() == "true"
    fmt = request.args.get("format","csv").lower()
    results = json.loads(row["results"] or "[]")
    if criteria_only:
        results = [r for r in results if r.get("meets_all_criteria")]
    params = json.loads(row["params"] or "{}")
    ts  = datetime.utcnow().strftime("%Y%m%d_%H%M")
    sic = params.get("sic_code","SIC")
    pc  = params.get("postcode","").replace(" ","")
    if fmt == "excel":
        data     = _to_excel(results, params)
        filename = f"AcquireFlow_{sic}_{pc}_{ts}.xlsx"
        mime     = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        data     = _to_csv(results)
        filename = f"AcquireFlow_{sic}_{pc}_{ts}.csv"
        mime     = "text/csv"
    return Response(data, mimetype=mime,
                    headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@app.delete("/api/cache")
def clear_cache():
    with _conn() as c:
        count = c.execute("SELECT COUNT(*) FROM company_cache").fetchone()[0]
    _write("DELETE FROM company_cache")
    return jsonify({"cleared": count, "message": f"Cleared {count} cached companies"})


@app.get("/api/health")
def health():
    key = os.getenv("COMPANIES_HOUSE_API_KEY","")
    return jsonify({"status":"ok",
                    "api_key_configured": bool(key) and key != "your_api_key_here",
                    "version":"1.0.0"})


# ── Export helpers ───────────────────────────────────────────────────────────────

def _flatten(r, idx):
    dirs   = r.get("directors",[])
    active = [d for d in dirs if d.get("is_active")]
    qual   = [d for d in active if d.get("meets_age_criteria")]
    dir_ages = "; ".join(
        f"{d['name']} (~{d['estimated_age']})" if d.get("estimated_age") else d["name"]
        for d in active)
    return {
        "#": idx,
        "Company Name":              r.get("company_name",""),
        "Company Number":            r.get("company_number",""),
        "Status":                    r.get("company_status",""),
        "Registered Address":        r.get("registered_address",""),
        "Postcode":                  r.get("postcode",""),
        "Distance (miles)":          round(r["distance_miles"],2) if r.get("distance_miles") is not None else "",
        "Incorporation Date":        r.get("incorporation_date",""),
        "Company Age (years)":       round(r["company_age_years"],1) if r.get("company_age_years") is not None else "",
        "SIC Codes":                 ", ".join(r.get("sic_codes",[])),
        "Active Directors":          r.get("active_director_count",0),
        "Qualifying Directors":      r.get("qualifying_director_count",0),
        "Director Names & Ages":     dir_ages,
        "Qualifying Director Names": "; ".join(d["name"] for d in qual),
        "Director Criteria Met":     "Yes" if r.get("director_criteria_met") else "No",
        "Company Age Criteria Met":  "Yes" if r.get("company_age_criteria_met") else "No",
        "All Criteria Met":          "Yes" if r.get("meets_all_criteria") else "No",
        "Buy Score (1-10)":          r.get("score",0),
    }


def _to_csv(results):
    if not results:
        return b""
    rows = [_flatten(r,i+1) for i,r in enumerate(results)]
    buf  = io.StringIO()
    w    = csv.DictWriter(buf, fieldnames=rows[0].keys())
    w.writeheader(); w.writerows(rows)
    return buf.getvalue().encode("utf-8-sig")


def _to_excel(results, search_params):
    wb = Workbook(); ws = wb.active; ws.title = "Results"
    if not results:
        ws["A1"] = "No results."; buf = io.BytesIO(); wb.save(buf); return buf.getvalue()

    rows    = [_flatten(r,i+1) for i,r in enumerate(results)]
    headers = list(rows[0].keys())
    hfont   = Font(bold=True, color="FFFFFF", size=11)
    hfill   = PatternFill("solid", fgColor="1E3A5F")
    qfill   = PatternFill("solid", fgColor="D4EDDA")
    thin    = Side(style="thin", color="CCCCCC")
    brd     = Border(left=thin, right=thin, top=thin, bottom=thin)
    ctr     = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wrp     = Alignment(horizontal="left",   vertical="top",    wrap_text=True)

    ws.row_dimensions[1].height = 30
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = hfont; cell.fill = hfill; cell.alignment = ctr; cell.border = brd

    for ri, row in enumerate(rows, 2):
        ws.row_dimensions[ri].height = 32
        is_q = row.get("All Criteria Met") == "Yes"
        for ci, key in enumerate(headers, 1):
            cell = ws.cell(row=ri, column=ci, value=row[key])
            cell.border = brd; cell.alignment = wrp
            if is_q: cell.fill = qfill

    sc = headers.index("Buy Score (1-10)") + 1
    for ri, row in enumerate(rows, 2):
        s = row.get("Buy Score (1-10)") or 0
        ws.cell(ri, sc).font = Font(bold=True,
            color="155724" if s >= 7 else "856404" if s >= 4 else "721C24")

    widths = {"#":5,"Company Name":35,"Company Number":16,"Status":10,
              "Registered Address":40,"Postcode":12,"Distance (miles)":14,
              "Incorporation Date":16,"Company Age (years)":16,"SIC Codes":14,
              "Active Directors":14,"Qualifying Directors":18,
              "Director Names & Ages":45,"Qualifying Director Names":35,
              "Director Criteria Met":18,"Company Age Criteria Met":20,
              "All Criteria Met":14,"Buy Score (1-10)":14}
    for col, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(col)].width = widths.get(h, 15)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    ws2 = wb.create_sheet("Search Summary")
    ws2["A1"] = "AcquireFlow – Search Summary"
    ws2["A1"].font = Font(bold=True, size=14)
    summary = [("SIC Code", search_params.get("sic_code","")),
               ("Postcode / Town", search_params.get("postcode","")),
               ("Radius (miles)", search_params.get("radius_miles","")),
               ("Min Director Age", search_params.get("min_director_age","")),
               ("Min Company Age (years)", search_params.get("min_company_age","")),
               ("Total Companies", len(results)),
               ("Criteria Met", sum(1 for r in rows if r["All Criteria Met"]=="Yes")),
               ("Export Date", datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC"))]
    for i, (k, v) in enumerate(summary, 3):
        ws2.cell(i, 1, k).font = Font(bold=True)
        ws2.cell(i, 2, str(v))
    ws2.column_dimensions["A"].width = 30
    ws2.column_dimensions["B"].width = 30

    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()


# ── Start ────────────────────────────────────────────────────────────────────────

if __name__ == "__main__":
    init_db()
    app.run(host="127.0.0.1", port=8000, debug=False, threaded=True)
